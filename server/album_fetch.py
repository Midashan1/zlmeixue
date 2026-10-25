#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
拉取公众号「合集」的全部文章（标题 / 发表时间 / 永久链接），可追加到已有 Excel 的新工作表。

为什么走这条路：
  公众号「合集」(album) 的 getalbum 接口是**公开**的，不需要登录、不需要 key。
  而 mp/profile_ext 历史消息接口需要微信会话（无 key 返回 no session）。
  所以「补齐永久链接」这件事可以全免费自动化。

拿不到的：阅读数 / 点赞数 / 转发数 —— getappmsgext 无会话时返回 ret:0 但 appmsgstat 为空，
  必须靠后台导出或带 key 的会话，不在本脚本范围内。

用法：
    # 一个文章链接
    python album_fetch.py --article https://mp.weixin.qq.com/s/xxxxxxxx
    # 多个合集各给一篇（自动去重合并）
    python album_fetch.py --article <url1> <url2> --output out/all_articles.xlsx
    # 直接给合集页地址
    python album_fetch.py --album-url "https://mp.weixin.qq.com/mp/appmsgalbum?__biz=xxx&album_id=yyy"
    # 追加到已有 Excel 的新工作表
    python album_fetch.py --album-url <url1> <url2> --append-to out/all_articles.xlsx \
        --sheet "二号·文章清单" --stats-sheet "二号·合集统计"
"""
from __future__ import annotations

import argparse
import json
import re
import time
import urllib.parse
import urllib.request
from datetime import datetime, timezone, timedelta
from pathlib import Path

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")
TZ = timezone(timedelta(hours=8))
HEADER_FILL = "2F5597"


# ---------------------------------------------------------------- 抓取
def fetch(url: str, referer: str = "https://mp.weixin.qq.com/") -> str:
    req = urllib.request.Request(url, headers={
        "User-Agent": UA, "Referer": referer, "Accept-Language": "zh-CN,zh;q=0.9",
    })
    with urllib.request.urlopen(req, timeout=25) as r:
        return r.read().decode("utf-8", "ignore")


def parse_album_url(url: str) -> tuple[str, str]:
    q = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
    biz = (q.get("__biz") or [""])[0]
    aid = (q.get("album_id") or [""])[0]
    if not biz or not aid:
        raise SystemExit(f"× 合集地址里缺 __biz 或 album_id：{url}")
    return biz, aid


def discover(article_url: str) -> dict:
    """从文章页提取 __biz、album_id 列表、公众号昵称、文章标题。带重试（微信偶发风控页）。"""
    last = ""
    for attempt in range(4):
        html = fetch(article_url)
        biz = re.search(r'var\s+biz\s*=\s*"([^"]+)"', html)
        if biz:
            albums = re.findall(r"albumId(?:Str)?\s*[:=]\s*'(\d{6,25})'", html)
            nick = re.search(r'var\s+nickname\s*=\s*"([^"]*)"', html)
            title = re.search(r'<meta\s+property="og:title"\s+content="([^"]*)"', html)
            return {
                "biz": biz.group(1),
                "albums": list(dict.fromkeys(albums)),
                "nick": nick.group(1) if nick else "",
                "title": title.group(1) if title else "",
            }
        last = html[:120].replace("\n", " ")
        print(f"    · 第 {attempt + 1} 次没拿到正文，{1.5 * (attempt + 1):.1f}s 后重试")
        time.sleep(1.5 * (attempt + 1))
    raise SystemExit(f"× 4 次都没找到 __biz（可能被风控）：{article_url}\n  页面片段：{last}")


def album_page(biz: str, album_id: str, begin_msgid: str = "", begin_itemidx: str = "") -> dict:
    q = urllib.parse.urlencode({
        "action": "getalbum", "__biz": biz, "album_id": album_id, "count": "20",
        "begin_msgid": begin_msgid, "begin_itemidx": begin_itemidx, "f": "json",
    })
    ref = f"https://mp.weixin.qq.com/mp/appmsgalbum?action=getalbum&album_id={album_id}"
    txt = fetch(f"https://mp.weixin.qq.com/mp/appmsgalbum?{q}", referer=ref)
    try:
        data = json.loads(txt)
    except json.JSONDecodeError:
        raise SystemExit(f"× 返回不是 JSON（可能被风控）：{txt[:200]}")
    if data.get("base_resp", {}).get("ret") not in (0, "0", None):
        raise SystemExit(f"× 接口报错：{data.get('base_resp')}")
    return data.get("getalbum_resp") or {}


def walk_album(biz: str, album_id: str) -> tuple[list[dict], dict]:
    rows, meta = [], {}
    begin_msgid = begin_itemidx = ""
    page = 1
    while True:
        resp = album_page(biz, album_id, begin_msgid, begin_itemidx)
        if page == 1:
            meta = resp.get("base_info") or {}
        arts = resp.get("article_list") or []
        print(f"    第 {page} 页：{len(arts)} 条")
        for a in arts:
            if not a.get("title"):
                continue
            rows.append({
                "标题": a.get("title", "").strip(),
                "发表时间": datetime.fromtimestamp(int(a["create_time"]), TZ).strftime("%Y-%m-%d %H:%M")
                if a.get("create_time") else "",
                "链接": (a.get("url") or "").replace("#rd", ""),
                "msgid": a.get("msgid", ""),
                "所属合集": meta.get("title", "") if meta else "",
            })
        if str(resp.get("continue_flag", "0")) != "1" or not arts:
            break
        begin_msgid = arts[-1].get("msgid", "")
        begin_itemidx = arts[-1].get("itemidx", "1")
        page += 1
        if page > 60:
            print("    ! 触达页数上限，停止")
            break
    return rows, meta


# ---------------------------------------------------------------- 写盘
def _write_sheet(wb, sheet_name: str, df):
    """在 workbook 里新建（或覆盖）一个工作表，写入 df 并套用统一样式。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    cols = list(df.columns)
    ws.append(cols)
    for rec in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and v != v) else v for v in rec])

    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 22

    widths = {"序号": 6, "标题": 56, "发表时间": 18, "所属合集": 16,
              "链接": 72, "msgid": 15, "篇数": 10, "最早": 18, "最新": 18}
    for i, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 14)

    if "发表时间" in cols:
        idx = cols.index("发表时间") + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=idx)
            if hasattr(cell.value, "strftime"):
                cell.number_format = "yyyy-mm-dd hh:mm"
    if "链接" in cols:
        idx = cols.index("链接") + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=idx)
            v = cell.value
            if isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(color="0563C1", underline="single")
    return ws


def build_frames(rows: list[dict]):
    import pandas as pd
    df = pd.DataFrame(rows).sort_values("发表时间", ascending=False).reset_index(drop=True)
    df.insert(0, "序号", range(1, len(df) + 1))
    df = df[["序号", "标题", "发表时间", "所属合集", "链接", "msgid"]]
    stats = (df.groupby("所属合集")
               .agg(篇数=("标题", "size"), 最早=("发表时间", "min"), 最新=("发表时间", "max"))
               .reset_index().sort_values("篇数", ascending=False))
    return df, stats


def write_new(path: Path, df, stats) -> None:
    import pandas as pd
    from openpyxl import Workbook
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "文章清单", df)
    _write_sheet(wb, "合集统计", stats)
    wb.save(path)


def append_to(path: Path, df, stats, sheet: str, stats_sheet: str) -> None:
    from openpyxl import load_workbook
    wb = load_workbook(path)
    _write_sheet(wb, sheet, df)
    if stats_sheet:
        _write_sheet(wb, stats_sheet, stats)
    wb.save(path)


# ---------------------------------------------------------------- 主流程
def main() -> None:
    ap = argparse.ArgumentParser(description="拉取公众号合集文章清单")
    ap.add_argument("--article", nargs="+", help="一篇或多篇该号文章链接（每个合集给一篇）")
    ap.add_argument("--album-url", nargs="+", help="合集页地址（含 __biz 与 album_id）")
    ap.add_argument("--biz", help="__biz（与 --album-id 一起用）")
    ap.add_argument("--album-id", nargs="+", help="合集 id")
    ap.add_argument("--output", default="out/album_articles.xlsx")
    ap.add_argument("--append-to", help="已有 xlsx 路径：结果写入它的新工作表")
    ap.add_argument("--sheet", default="文章清单", help="文章清单工作表名")
    ap.add_argument("--stats-sheet", default="", help="合集统计工作表名（留空则不写）")
    args = ap.parse_args()

    biz, nick, album_ids = None, "", []

    if args.album_url:
        for u in args.album_url:
            b, a = parse_album_url(u)
            biz = biz or b
            if b != biz:
                raise SystemExit(f"× 这批地址 __biz 不一致：{b} ≠ {biz}")
            if a in album_ids:
                print(f"  · 合集 {a} 重复出现，跳过")
                continue
            album_ids.append(a)
        print(f"合集地址 {len(args.album_url)} 条 → 去重后 {len(album_ids)} 个合集")

    if args.article:
        seen = set(album_ids)
        for i, url in enumerate(args.article, 1):
            print(f"[{i}/{len(args.article)}] 解析 {url[-22:]}")
            info = discover(url)
            biz = biz or info["biz"]
            nick = nick or info["nick"]
            if not info["albums"]:
                print("    ! 这篇文章没有挂合集，跳过")
                continue
            for aid in info["albums"]:
                if aid not in seen:
                    seen.add(aid)
                    album_ids.append(aid)
                    print(f"    → 新合集 {aid}（种子文章：{info['title'][:24]}）")
                else:
                    print(f"    → 合集 {aid} 已收录，跳过")

    if args.biz and args.album_id:
        biz = biz or args.biz
        album_ids.extend(a for a in args.album_id if a not in album_ids)

    if not biz or not album_ids:
        raise SystemExit("× 需要 --article / --album-url，或 --biz 配 --album-id")

    all_rows, names = [], []
    for n, aid in enumerate(album_ids, 1):
        print(f"\n[{n}/{len(album_ids)}] 拉取合集 {aid} …")
        rows, meta = walk_album(biz, aid)
        name = meta.get("title", "") or f"album_{aid}"
        names.append(name)
        if not nick:
            nick = meta.get("nickname", "")
        print(f"    合集《{name}》计入 {len(rows)} 篇")
        all_rows.extend(rows)

    if not all_rows:
        raise SystemExit("× 没取到文章，合集可能已删除或未公开")

    seen, uniq, dup = set(), [], 0
    for r in all_rows:
        k = r["msgid"] or r["链接"]
        if k in seen:
            dup += 1
            continue
        seen.add(k)
        uniq.append(r)
    if dup:
        print(f"\n跨合集重复 {dup} 篇，已去重")

    df, stats = build_frames(uniq)

    if args.append_to:
        target = Path(args.append_to)
        if not target.exists():
            raise SystemExit(f"× 目标文件不存在：{target}")
        append_to(target, df, stats, args.sheet, args.stats_sheet)
        print(f"\n已写入 {target.resolve()} 的工作表「{args.sheet}」"
              + (f" + 「{args.stats_sheet}」" if args.stats_sheet else ""))
    else:
        write_new(Path(args.output), df, stats)
        print(f"\n已写入 {Path(args.output).resolve()}")

    print(f"公众号：{nick or '（未取到）'}｜合集 {len(names)} 个：{'、'.join(names)}")
    print(f"合并去重后 {len(df)} 篇，跨度 {df['发表时间'].min()} ~ {df['发表时间'].max()}")
    print("提醒：本清单含 标题/时间/合集/永久链接；阅读、点赞、转发需后台导出补。")


if __name__ == "__main__":
    main()

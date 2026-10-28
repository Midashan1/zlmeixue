#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
一条命令刷新整个文章库：按 accounts.json 里的合集清单抓取 → 输出 articles.json（+ 可选 xlsx）。

给定时任务用（每天跑一次即可）：
    python refresh.py                 # 只更新 articles.json
    python refresh.py --xlsx          # 顺带更新 out/all_articles.xlsx
    python refresh.py --build         # 更新完再重新生成 HTML

输出：
    deploy/articles.json   页面运行时读取的数据（更新只需替换这一个文件）
    out/all_articles.xlsx  人工查看用（--xlsx）
    out/文章库.html        离线单文件版（--build）
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

import album_fetch as af

BASE = Path(__file__).parent
CONFIG = BASE / "accounts.json"


def load_config() -> list[dict]:
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
    accts = cfg.get("accounts") or []
    if not accts:
        raise SystemExit(f"× {CONFIG} 里没有 accounts")
    for a in accts:
        if not a.get("biz") or not a.get("albums"):
            raise SystemExit(f"× 账号 {a.get('name')} 缺 biz 或 albums")
    return accts


def refresh(quiet: bool = False) -> tuple[list[dict], list[dict]]:
    """返回 (去重后的文章行, 失败合集列表)。单个合集失败不影响整批。"""
    rows: list[dict] = []
    failures: list[dict] = []
    for acct in load_config():
        name, biz = acct["name"], acct["biz"]
        if not quiet:
            print(f"\n=== {name}（{len(acct['albums'])} 个合集）")
        for i, aid in enumerate(acct["albums"], 1):
            if not quiet:
                print(f"  [{i}/{len(acct['albums'])}] 合集 {aid}")
            try:
                arts, meta = af.walk_album(biz, aid)
            except (SystemExit, Exception) as exc:   # 单个合集失败不拖垮整批
                msg = str(exc) or exc.__class__.__name__
                print(f"    ! 跳过：{msg}")
                failures.append({"acct": name, "album_id": aid, "error": msg[:200]})
                continue
            album_name = meta.get("title") or f"album_{aid}"
            for r in arts:
                rows.append({
                    "title": r["标题"],
                    "ts": r["发表时间"],
                    "album": album_name,
                    "acct": name,
                    "url": r["链接"],
                })
            if not quiet:
                print(f"    → 《{album_name}》{len(arts)} 篇")
            time.sleep(0.6)                    # 温和抓取，避免风控

    # 去重（同一篇可能挂在多个合集）
    seen, uniq = set(), []
    for r in rows:
        key = r["url"].split("&chksm=")[0]
        if key in seen:
            continue
        seen.add(key)
        uniq.append(r)
    uniq.sort(key=lambda r: r["ts"], reverse=True)
    return uniq, failures


def merge_previous(rows: list[dict], prev: list[dict], quiet: bool = False) -> tuple[list[dict], list[str]]:
    """某个合集这次抓到的条数明显变少（多半是被风控/接口异常），用上次的数据补齐，
    避免"静默丢文章"。返回 (合并后的行, 被回填的合集名)。"""
    from collections import defaultdict

    def group(data):
        g = defaultdict(list)
        for r in data:
            g[(r.get("acct", ""), r.get("album", ""))].append(r)
        return g

    prev_g, new_g = group(prev), group(rows)
    restored = []
    for key, plist in prev_g.items():
        if not plist:
            continue
        n = len(new_g.get(key, []))
        if n >= len(plist) * 0.8:          # 抓全了或正常增减，不干预
            continue
        have = {r["url"].split("&chksm=")[0] for r in new_g.get(key, [])}
        add = [r for r in plist if r["url"].split("&chksm=")[0] not in have]
        if add:
            rows.extend(add)
            restored.append(f"{key[0]}·{key[1]}（补回 {len(add)} 篇）")
            if not quiet:
                print(f"    · 回填《{key[1]}》{len(add)} 篇（本次只抓到 {n}/{len(plist)}）")

    seen, out = set(), []
    for r in rows:
        k = r["url"].split("&chksm=")[0]
        if k in seen:
            continue
        seen.add(k)
        out.append(r)
    out.sort(key=lambda r: r["ts"], reverse=True)
    return out, restored


def write_json(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
    path.write_text(payload, encoding="utf-8")


def write_xlsx(rows: list[dict], path: Path) -> None:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(path) if path.exists() else None
    if wb is None:
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)

    # 分表顺序必须跟 accounts.json 一致：build_library.py 约定「文章清单」= 一号账号。
    # 若按 rows 出现顺序建表，哪号最新发过文章就会抢到「文章清单」，导致两个账号标签互换。
    by_acct: dict[str, list[dict]] = {a["name"]: [] for a in load_config()}
    for r in rows:
        by_acct.setdefault(r["acct"], []).append(r)
    by_acct = {k: v for k, v in by_acct.items() if v}   # 空账号不出空表

    SHEETS = ["文章清单", "二号·文章清单", "三号·文章清单"]
    for idx, (name, items) in enumerate(by_acct.items()):
        df = pd.DataFrame(items).rename(columns={
            "title": "标题", "ts": "发表时间", "album": "所属合集", "url": "链接"
        })
        df = df[["标题", "发表时间", "所属合集", "链接"]]
        df.insert(0, "序号", range(1, len(df) + 1))
        sheet = SHEETS[idx] if idx < len(SHEETS) else f"{name}·文章清单"
        if sheet in wb.sheetnames:
            del wb[sheet]
        ws = wb.create_sheet(sheet)
        ws.append(list(df.columns))
        for rec in df.itertuples(index=False):
            ws.append(list(rec))
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor="2F5597")
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, w in {"A": 6, "B": 56, "C": 18, "D": 16, "E": 72}.items():
            ws.column_dimensions[col].width = w
        link_col = get_column_letter(5)
        for r in range(2, ws.max_row + 1):
            cell = ws[f"{link_col}{r}"]
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    wb.save(path)


def load_previous(path: Path) -> list[dict]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except Exception:
        return []


def write_status(path: Path, rows: list[dict], failures: list[dict], restored: list[str]) -> None:
    """写一份刷新状态，方便排查"为什么没更新"。"""
    from datetime import datetime, timezone, timedelta

    accts: dict[str, int] = {}
    for r in rows:
        accts[r["acct"]] = accts.get(r["acct"], 0) + 1
    status = {
        "refreshed_at": datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S"),
        "total": len(rows),
        "by_account": accts,
        "newest": {"ts": rows[0]["ts"], "title": rows[0]["title"]} if rows else None,
        "failed_albums": failures,
        "restored_albums": restored,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser(description="刷新文章库数据")
    ap.add_argument("--json", default=str(BASE / "deploy" / "articles.json"))
    ap.add_argument("--xlsx", action="store_true", help="同时更新 out/all_articles.xlsx")
    ap.add_argument("--build", action="store_true", help="同时重建离线 HTML")
    ap.add_argument("--force", action="store_true", help="篇数大幅缩水时也照写")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    out_json = Path(args.json)
    prev = load_previous(out_json)

    rows, failures = refresh(quiet=args.quiet)
    restored: list[str] = []
    if prev:
        rows, restored = merge_previous(rows, prev, quiet=args.quiet)

    if not rows:
        raise SystemExit("× 一篇都没抓到，检查网络或 accounts.json")

    # 防止一次异常把已有数据冲成半截
    if prev and len(rows) < len(prev) * 0.9 and not args.force:
        raise SystemExit(
            f"× 本次只抓到 {len(rows)} 篇，上次有 {len(prev)} 篇，疑似大面积失败；"
            f"已保留旧数据。确认无误可加 --force"
        )

    write_json(rows, out_json)
    # 状态文件放在 articles.json 旁边：线上即可直接访问 /zlmeixue/status.json 自查
    write_status(out_json.parent / "status.json", rows, failures, restored)

    accts = {}
    for r in rows:
        accts[r["acct"]] = accts.get(r["acct"], 0) + 1
    print(f"\n共 {len(rows)} 篇 → {out_json}")
    print("按账号：" + " / ".join(f"{k} {v} 篇" for k, v in accts.items()))
    print(f"最新一篇：{rows[0]['ts']}  {rows[0]['title'][:30]}")
    if failures:
        print(f"! 有 {len(failures)} 个合集抓取失败：" + "、".join(f["album_id"] for f in failures))
    if restored:
        print("· 已用旧数据回填：" + "；".join(restored))

    if args.xlsx:
        write_xlsx(rows, BASE / "out" / "all_articles.xlsx")
        print(f"已更新 {BASE / 'out' / 'all_articles.xlsx'}")

    if args.build:
        import subprocess
        subprocess.run([sys.executable, str(BASE / "build_library.py"),
                        "--input", str(BASE / "out" / "all_articles.xlsx"),
                        "--output", str(BASE / "out" / "文章库.html")], check=False)


if __name__ == "__main__":
    main()
